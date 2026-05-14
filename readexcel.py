"""Extract text from Excel (.xlsx / .xls) spreadsheets and save as Markdown."""

import re
import time
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None

INPUT_DIR = Path("input")
OUTPUT_DIR = Path("output_md")


def sanitize_filename(name: str) -> str:
    """Remove or replace characters unsafe for filenames."""
    name = re.sub(r"[\\/:*?\"<>|]", "_", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name or "untitled"


def cell_value(cell) -> str:
    """Convert a cell to string, handling None, numbers, dates."""
    if cell.value is None:
        return ""
    if isinstance(cell.value, float):
        if cell.value == int(cell.value):
            return str(int(cell.value))
        return str(cell.value)
    return str(cell.value)



def extract_xls_text(xls_path: Path) -> str:
    """Extract all text from an .xls file using xlrd."""
    wb = xlrd.open_workbook(str(xls_path))
    sections = []

    for sheet_idx in range(wb.nsheets):
        ws = wb.sheet_by_index(sheet_idx)
        sheet_name = ws.name
        lines = []

        if ws.ncols == 0:
            continue

        if sheet_name:
            sections.append(f"## Sheet: {sheet_name}")

        if ws.ncols <= 1:
            # Single column
            for row_idx in range(ws.nrows):
                val = str(ws.cell_value(row_idx, 0)).strip()
                if val:
                    lines.append(f"- {val}")
        else:
            # Markdown table
            for row_idx in range(ws.nrows):
                row_vals = [str(ws.cell_value(row_idx, c)).strip() for c in range(ws.ncols)]
                lines.append("| " + " | ".join(row_vals) + " |")
                if row_idx == 0 and ws.nrows > 1:
                    lines.append("| " + " | ".join(["---"] * ws.ncols) + " |")

        text = "\n".join(lines)
        if text.strip():
            sections.append(text)

    return "\n\n".join(sections)



def extract_xlsx_text(xlsx_path: Path) -> str:
    """Extract all text from an .xlsx file, sheet by sheet."""
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True, read_only=True)
    sections = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines = []
        max_col = 0
        all_rows = []

        # Read all rows to determine max column width
        for row in ws.iter_rows(values_only=False):
            values = [cell_value(cell) for cell in row]
            all_rows.append(values)
            # Find last non-empty column index
            non_empty = [i for i, v in enumerate(values) if v.strip()]
            if non_empty:
                max_col = max(max_col, non_empty[-1] + 1)

        if not all_rows:
            continue

        # Trim each row to max_col
        all_rows = [row[:max_col] for row in all_rows]

        # Detect if this is a table (first row looks like header)
        first_row = all_rows[0] if all_rows else []
        has_header = any(bool(v.strip()) for v in first_row)

        # Build markdown
        if sheet_name:
            sections.append(f"## Sheet: {sheet_name}")

        if max_col <= 1:
            # Single column → simple list
            for row in all_rows:
                val = row[0].strip() if row else ""
                if val:
                    lines.append(f"- {val}")
        else:
            # Multi-column → markdown table
            for idx, row in enumerate(all_rows):
                formatted = [v.strip() for v in row]
                lines.append("| " + " | ".join(formatted) + " |")
                if idx == 0 and has_header and len(all_rows) > 1:
                    # Separator row
                    lines.append("| " + " | ".join(["---"] * max_col) + " |")

        text = "\n".join(lines)
        if text.strip():
            sections.append(text)

    wb.close()
    return "\n\n".join(sections)


def main() -> None:
    if openpyxl is None:
        print("请安装 openpyxl: pip install openpyxl")
        return

    OUTPUT_DIR.mkdir(exist_ok=True)

    excel_files = sorted(
        [p for p in INPUT_DIR.iterdir() if p.suffix.lower() in (".xlsx", ".xls")]
    )
    if not excel_files:
        print("No .xlsx / .xls files found in input/.")
        return

    total_start = time.time()
    for src_path in excel_files:
        file_start = time.time()
        print(f"Processing: {src_path.name}")

        is_xls = src_path.suffix.lower() == ".xls"
        if is_xls:
            if xlrd is None:
                print(f"  请安装 xlrd: pip install xlrd")
                continue
            text = extract_xls_text(src_path)
        else:
            text = extract_xlsx_text(src_path)

        if not text.strip():
            print(f"  Warning: No text extracted from {src_path.name}")
            text = "[No extractable text content.]"

        md_name = sanitize_filename(src_path.stem + ".md")
        md_path = OUTPUT_DIR / md_name
        md_content = f"# {src_path.stem}\n\n{text}\n"
        md_path.write_text(md_content, encoding="utf-8")

        elapsed = time.time() - file_start
        print(f"  Saved ({elapsed:.1f}s): {md_path}")

    total_elapsed = time.time() - total_start
    print(f"\nDone. {len(excel_files)} file(s) converted in {total_elapsed:.1f}s -> {OUTPUT_DIR}/")


if __name__ == "__main__":
    main()
